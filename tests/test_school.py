from collections import Counter
from copy import deepcopy
from io import BytesIO
import pytest
from openpyxl import load_workbook
from wfla.intake import school_dataset, school_demo, add_combination, add_staff_pool, add_common_courses, choice_template, import_choices
from wfla.schema import active_slots, activity_students, slot
from wfla.validation import validate
from wfla.solver import solve
from wfla.audit import audit
from wfla.io import timetable_rows, schedule_xlsx
from wfla.web import create_app


@pytest.fixture(scope='module')
def flexible():
    d=school_demo(1);r=solve(d,30)
    assert r['status'] in ('OPTIMAL','FEASIBLE'),r
    assert audit(d,r)==[]
    return d,r


def test_actual_calendars_and_weekly_totals(flexible):
    d,r=flexible
    assert len(active_slots(d,11))==44 and len(active_slots(d,12))==39
    students={s['id']:s for s in d['students']}
    rows=timetable_rows(d,r)
    for sid,s in students.items():
        mine=[row for row in rows if sid in row['学生'].split(';')]
        assert sum(row['节数'] for row in mine)==(44 if s['grade']==11 else 39)
        counts=Counter()
        for row in mine:counts[row['课程']]+=row['节数']
        assert counts['TOK']==(4 if s['grade']==11 else 3)
        assert counts['自修']==(5 if s['grade']==11 else 4)
        assert counts['English B']==(5 if s['grade']==11 else 6)
        assert counts['待安排']==(0 if s['grade']==11 else 1)
        assert counts['体育']==(1 if s['grade']==11 else 0)
        assert all(slot(d,row['星期'],row['起始节']) in active_slots(d,s['grade']) for row in mine)


def test_common_courses_teachers_take_two_separate_classes(flexible):
    d,r=flexible
    for grade in (11,12):
        tok=[a for a in d['activities'] if a['grade']==grade and a['name']=='TOK']
        assert len(tok)==6 and sorted(Counter(a['teacher'] for a in tok).values())==[2,2,2]
        assert all(len({next(s['admin_class'] for s in d['students'] if s['id']==sid) for sid in activity_students(d,a)})==1 for a in tok)


@pytest.mark.parametrize('mutation',['grade_closed','study_hours','study_capacity','study_wrong_room','common_hours','common_teacher_overlap'])
def test_flexible_independent_audit(flexible,mutation):
    d,r=deepcopy(flexible)
    if mutation=='grade_closed':
        section=next(s for s in r['sections'] if '-12-' in s['course'])
        next(l for l in r['lessons'] if l['section']==section['id'])['start']=slot(d,5,8)
    if mutation=='study_hours':r['studies'].pop()
    if mutation=='study_capacity':
        l=r['studies'][0];next(x for x in d['rooms'] if x['id']==l['room'])['capacity']=len(l['students'])-1
    if mutation=='study_wrong_room':r['studies'][0]['room']='C201'
    if mutation=='common_hours':r['activities'].pop()
    if mutation=='common_teacher_overlap':
        a=r['activities'][0];b=next(l for l in r['activities'][1:] if l['activity']==a['activity']);b['start']=a['start']
    assert audit(d,r)


def test_import_simple_excel_and_statistics_are_atomic():
    d=school_dataset()
    wb=load_workbook(BytesIO(choice_template(d,'combinations')))
    courses=['MAT-11-HL','PHY-11-HL','CHE-11-HL','CHA-11-SL','ECO-11-SL','ENB-11-SL']
    wb['选课组合'].append([11,'11-1',3,*courses]);stream=BytesIO();wb.save(stream)
    result,count=import_choices(d,stream.getvalue(),'统计.xlsx')
    assert count==3 and len(result['students'])==3 and len(result['choices'])==18 and d['students']==[]
    assert all(s['_row']==2 for s in result['students'])
    wb['选课组合'].append([11,'11-2',3,*(['MAT-11-HL']*6)]);stream=BytesIO();wb.save(stream)
    with pytest.raises(ValueError,match='第3行'):import_choices(d,stream.getvalue(),'bad.xlsx')
    assert d['students']==[]


def test_resource_profile_and_staff_not_double_counted():
    d=school_dataset();rooms={r['id']:r for r in d['rooms']}
    assert 'A308' in rooms and 'A302' not in rooms and 'A305' not in rooms
    assert rooms['C201']['capacity']==44 and rooms['C301']['capacity']==22
    assert len([r for r in rooms if r.startswith(('A3','A4','A5'))])==24
    d=add_staff_pool(d,'数学','MAT',['HL','SL'],[11,12],2,24,['甲老师','乙老师'])
    assert len(d['teachers'])==2 and len(d['qualifications'])==8
    with pytest.raises(ValueError,match='同名'):add_staff_pool(d,'兼教','PHY',['HL'],[11],1,24,['甲老师'])
    d['rooms'].append(dict(id='A302',name='办公室',capacity=22,types='普通',preference=0))
    assert any('办公室' in e['message'] for e in validate(d,False))


def test_pooled_rooms_and_cohort_study_still_allow_splitting():
    d=school_dataset();d['config']['quick_feasible']=True
    selected=['MAT-11-HL','PHY-11-HL','CHE-11-HL','CHA-11-SL','ECO-11-SL','ENB-11-SL']
    d=add_combination(d,11,'11-1',8,selected)
    for family in ('MAT','PHY','CHE','CHA','ECO','ENB'):d=add_staff_pool(d,family,family,['HL','SL'],[11],1,24)
    d['teachers'].append(dict(id='CORE',name='虚构公共课教师',max_periods=24))
    d=add_common_courses(d,11,{name:['CORE'] for name in ['TOK','体育','会考地理','会考历史','会考政治','Career Navigation']})
    for rid in ('LAB1','LAB2'):d['rooms'].append(dict(id=rid,name=rid,capacity=15,types='普通;物理实验;化学实验',preference=2))
    for room in d['rooms']:
        if room['id'] in ('C301','C302'):room['capacity']=4
    for c in d['courses']:
        if c['id'] not in selected:continue
        c['max_sections']=1;sid='CLASS:'+c['id']
        d['sections'].append(dict(id=sid,course=c['id'],teacher='',room='',block=0))
        d['enrollments'].extend(dict(student=s['id'],section=sid) for s in d['students'])
    from wfla.schema import normalize
    d=normalize(d);r=solve(d,30)
    assert r['status'] in ('OPTIMAL','FEASIBLE'),r
    assert audit(d,r)==[]
    assert len(r['studies'])==10
    assert all(len(s['students'])==4 for s in r['studies'])


def test_weekly_budget_mismatch_is_blocked():
    d=school_demo(1);next(r for r in d['study_rules'] if r['grade']==12)['unassigned']=0
    assert any('38节' in e['message'] for e in validate(d))
    assert solve(d,1)['status']=='INVALID_DATA'


def test_shared_study_room_capacity_not_per_grade(flexible):
    d,r=flexible
    for session in r['studies']:
        capacity=next(room['capacity'] for room in d['rooms'] if room['id']==session['room'])
        assert len(session['students'])<=capacity
    assert len({(s['room'],s['start']) for s in r['studies']})==len(r['studies'])


def test_export_new_row_types_keep_column_order(flexible):
    d,r=flexible;wb=load_workbook(BytesIO(schedule_xlsx(d,r)))
    values=list(wb['课次明细'].values);header=values[0]
    rows=[dict(zip(header,row)) for row in values[1:]]
    assert any(row['课程']=='TOK' and row['课型']=='公共课' for row in rows)
    assert any(row['课程']=='自修' and row['教室'] in ('C301','C302') for row in rows)
    assert any(row['课程']=='待安排' for row in rows)


def test_new_pages_and_intake_flow(tmp_path,flexible):
    app=create_app(tmp_path/'web.sqlite3',testing=True);client=app.test_client();client.get('/')
    with client.session_transaction() as s:token=s['csrf'];v=s['dataset_id']
    for path in ('/','/intake','/staff','/common','/rules','/manage'):
        response=client.get(path);assert response.status_code==200,(path,response.text)
    assert '录入选课' in client.get('/').text
    payload={'csrf':token,'v':v,'grade':11,'admin_class':'11-1','count':2,'action':'manual'}
    courses=['MAT-11-HL','PHY-11-HL','CHE-11-HL','CHA-11-SL','ECO-11-SL','ENB-11-SL']
    payload.update({f'choice{i+1}':c for i,c in enumerate(courses)})
    assert client.post('/intake',data=payload).status_code==302
    store=app.extensions['store']
    with client.session_transaction() as s:v=s['dataset_id']
    assert len(store.get(v)['students'])==2
    assert '暂不能排课' in client.get('/validation').text
    d,r=flexible;v=store.save(d);run=store.new_run(v);store.finish(run,r)
    for kind in ('student','teacher','room','section'):
        assert client.get(f'/run/{run}/view?kind={kind}').status_code==200
    sid=next(s['id'] for s in d['students'] if s['grade']==12)
    html=client.get(f'/run/{run}/view?kind=student&id={sid}').text
    assert '待安排' in html and '自修' in html and 'closed' in html
