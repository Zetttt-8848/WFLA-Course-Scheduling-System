from copy import deepcopy
import io
import json
import pytest
from openpyxl import load_workbook
from wfla.demo import demo_dataset
from wfla.schema import normalize
from wfla.validation import validate
from wfla.solver import solve
from wfla.audit import audit
from wfla.io import workbook_bytes,import_bytes,timetable_rows,schedule_xlsx,csv_bytes
from wfla.web import create_app


@pytest.fixture(scope="module")
def solved():
    data=demo_dataset(12)
    result=solve(data,30)
    assert result["status"] in ("OPTIMAL","FEASIBLE"),result
    assert not result["audit"]
    return data,result


def test_demo_rules_and_six_blocks(solved):
    data,result=solved
    assert not validate(data)
    assert not audit(data,result)
    for student in data["students"]:
        assert sorted(s["block"] for s in result["sections"] if student["id"] in s["students"])==[1,2,3,4,5,6]


def test_auto_sections_blocks_and_membership():
    d=demo_dataset(3,(11,))
    d["sections"]=[];d["enrollments"]=[];d["blocks"]=[]
    for c in d["courses"]:c["max_sections"]=1
    r=solve(d,30)
    assert r["status"] in ("OPTIMAL","FEASIBLE"),r
    assert not audit(d,r)


@pytest.mark.parametrize("mutation",["missing_choice","duplicate_choice","unapproved4","revoked4","wrong_term4","grade10_hl","grade10_english","teacher_missing","room_missing"])
def test_illegal_inputs(mutation):
    d=demo_dataset(12)
    if mutation=="missing_choice":d["choices"].pop()
    if mutation=="duplicate_choice":d["choices"].append(deepcopy(d["choices"][0]))
    if mutation=="unapproved4":d["approvals"]=[]
    if mutation=="revoked4":d["approvals"][0]["status"]="revoked"
    if mutation=="wrong_term4":d["approvals"][0]["term"]="OLD"
    if mutation=="grade10_hl":next(c for c in d["courses"] if c["grade"]==10)["level"]="HL"
    if mutation=="grade10_english":
        c=next(c for c in d["courses"] if c["grade"]==10);c["family"]="ENGLISH_A"
    if mutation=="teacher_missing":d["qualifications"]=[]
    if mutation=="room_missing":d["rooms"]=[r for r in d["rooms"] if not r["id"].startswith("LAB")]
    assert validate(d)
    assert solve(d,3)["status"]=="INVALID_DATA"


def test_fixed_events_conflict_caught_at_import():
    d=demo_dataset(12)
    e=deepcopy(d["events"][0]);e["id"]="CLASH";d["events"].append(e)
    assert any("重叠" in i["message"] for i in validate(d))


def test_teacher_unavailable_diagnosis():
    d=demo_dataset(3,(11,));course=d["courses"][0]["id"]
    targets={q["teacher"] for q in d["qualifications"] if q["course"]==course}
    for tid in targets:
        for day in range(1,6):
            for period in range(1,9):d["unavailable"].append(dict(kind="teacher",resource=tid,day=day,period=period,reason="测试禁排"))
    assert not validate(d)
    result=solve(d,10)
    assert result["status"]=="INFEASIBLE"
    assert "候选班数" in result["message"] and result["suggestions"]


@pytest.mark.parametrize("mutation",["room_type","room_capacity","missing_hours","bad_double","block_change","student_duplicate","teacher_overlap","room_unavailable","teacher_load"])
def test_independent_audit_detects_mutations(solved,mutation):
    d,r=deepcopy(solved)
    if mutation=="room_type":
        l=next(l for l in r["lessons"] if l["kind"]=="lab");l["room"]="R01"
    if mutation=="room_capacity":
        l=r["lessons"][0];next(room for room in d["rooms"] if room["id"]==l["room"])["capacity"]=0
    if mutation=="missing_hours":r["lessons"].pop()
    if mutation=="bad_double":
        l=next(l for l in r["lessons"] if l["duration"]==2);l["start"]=3
    if mutation=="block_change":r["sections"][0]["block"]=7
    if mutation=="student_duplicate":r["sections"][0]["students"].append(r["sections"][0]["students"][0])
    if mutation=="teacher_overlap":
        l1,l2=r["lessons"][:2];l2["start"]=l1["start"]
    if mutation=="room_unavailable":
        l=r["lessons"][0];d["unavailable"].append(dict(kind="room",resource=l["room"],day=l["start"]//8+1,period=l["start"]%8+1,reason="测试"))
    if mutation=="teacher_load":
        t=r["sections"][0]["teacher"];next(ti for ti in d["teachers"] if ti["id"]==t)["max_periods"]=0
    assert audit(d,r)


def test_excel_roundtrip_and_exports(solved):
    d,r=solved
    restored=import_bytes(workbook_bytes(d),"input.xlsx")
    assert restored==d
    out=schedule_xlsx(d,r)
    wb=load_workbook(io.BytesIO(out))
    assert {"学生视图","教师视图","教室视图","教学班视图"}<=set(wb.sheetnames)
    assert wb["课次明细"].max_row==len(timetable_rows(d,r))+1
    assert csv_bytes([{"姓名":"=cmd()"}]).decode("utf-8-sig").find("'=cmd()")>=0


def test_excel_formula_rejected():
    wb=load_workbook(io.BytesIO(workbook_bytes(demo_dataset(3,(11,)))))
    wb["学生"]["B2"]="=1+1"
    out=io.BytesIO();wb.save(out)
    with pytest.raises(ValueError,match="公式"):import_bytes(out.getvalue(),"input.xlsx")


def test_malformed_settings_report_error():
    d=demo_dataset(3);d["config"]["periods_per_day"]="bad"
    assert validate(d)[0]["table"]=="config"


def test_web_flow_versions_views_and_download(tmp_path,solved):
    app=create_app(tmp_path/"db.sqlite3",testing=True);client=app.test_client()
    assert client.get("/").status_code==200
    with client.session_transaction() as s:token=s["csrf"]
    assert client.post("/demo",data={"csrf":token,"size":"12"}).status_code==302
    assert "基础数据校验通过" in client.get("/validation").text
    store=app.extensions["store"]
    with client.session_transaction() as s:identifier=s["dataset_id"]
    run_id=store.new_run(identifier);store.finish(run_id,solved[1])
    assert client.get(f"/run/{run_id}").status_code==200
    for kind in ("student","teacher","room","section"):
        response=client.get(f"/run/{run_id}/view?kind={kind}")
        assert response.status_code==200
        assert "课表" in response.text
    for fmt in ("xlsx","csv","json"):
        assert client.get(f"/run/{run_id}/export/{fmt}").status_code==200
    old=store.get(identifier)
    assert client.post("/data/students",data={"csrf":token,"v":identifier,"action":"delete","row":0}).status_code==302
    assert store.get(identifier)==old
    assert "暂不能排课" in client.get("/validation").text
    assert client.post("/solve",data={"csrf":token,"seconds":5}).location.endswith("/validation")
    assert client.post("/demo",data={"size":12}).status_code==400
    assert client.get("/download/project.zip").status_code==200


def test_manual_invalid_move_preserves_original(tmp_path,solved):
    d,r=solved;app=create_app(tmp_path/"db.sqlite3",testing=True);client=app.test_client();client.get("/")
    store=app.extensions["store"];v=store.save(d);run=store.new_run(v);store.finish(run,r)
    with client.session_transaction() as s:token=s["csrf"]
    lesson=next(l for l in r["lessons"] if l["kind"]=="lab")
    response=client.post(f"/run/{run}/adjust",data={"csrf":token,"lesson":lesson["id"],"day":1,"period":1,"room":"R01"},follow_redirects=True)
    assert "调整未保存" in response.text
    assert store.run(run)["result"]==r
