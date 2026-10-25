"""School setup and simple, reversible intake; no actual student names bundled."""
from copy import deepcopy
import io
from openpyxl import Workbook, load_workbook
from .schema import empty_dataset, normalize

CATALOG = {
    'CHA':('Chinese A',1,'zh',''), 'ENB':('English B',2,'en',''),
    'MAT':('MAA',5,'',''), 'PHY':('Physics',4,'','物理实验'),
    'CHE':('Chemistry',4,'','化学实验'), 'BIO':('Biology',4,'','生物实验'),
    'ECO':('Economics',3,'',''), 'HIS':('History',3,'',''),
}


def school_dataset():
    d = empty_dataset()
    d['config'].update(title='新学期',term='新学期',days=5,periods_per_day=9,minutes_per_period=40,
                       scheduling_mode='flexible',school_profile=True)
    for grade in (11,12):
        for day in range(1,6):
            d['calendar'].append(dict(grade=grade,day=day,periods=(9 if day<5 else 8) if grade==11 else (8 if day<5 else 7)))
        d['study_rules'].append(dict(grade=grade,periods=5 if grade==11 else 4,rooms='C301;C302',unassigned=0 if grade==11 else 1))
        for family,(name,group,language,lab) in CATALOG.items():
            for level in ('HL','SL'):
                d['courses'].append(dict(id=f'{family}-{grade}-{level}',name=name,family=family,group=group,language=language,
                    grade=grade,level=level,periods=6 if grade==12 and family=='ENB' else 5,
                    lab_periods=2 if lab else 0,lab_doubles=1 if lab else 0,room_type='普通',lab_type=lab,capacity=15 if lab else 22))
    for floor in range(2,6):
        for number in (1,3,4,6,7,8,9,10):
            rid=f'A{floor}{number:02}'
            d['rooms'].append(dict(id=rid,name=rid,capacity=22,types='普通',preference=0 if floor>=3 else 1))
    for rid,cap in [('C201',44),('C202',44),('C301',22),('C302',22)]:
        d['rooms'].append(dict(id=rid,name=rid,capacity=cap,types='自修',preference=0))
    d['rooms'].append(dict(id='PLAYGROUND',name='操场',capacity=22,types='操场',preference=0))
    return normalize(d)


def add_combination(data, grade, admin_class, count, choices, roster=None):
    if grade not in (10,11,12) or not 1<=count<=400: raise ValueError('年级无效或人数不在1—400之间')
    if not admin_class.strip(): raise ValueError('请填写行政班；公共课按行政班安排')
    if len(choices)!=6 or len(set(choices))!=6: raise ValueError('请选择六门不同的课程')
    catalog={c['id']:c for c in data['courses']}
    if any(c not in catalog or catalog[c]['grade']!=grade for c in choices): raise ValueError('所选课程不属于该年级')
    from .validation import valid_combination
    selected=[catalog[c] for c in choices]
    if not valid_combination(selected): raise ValueError('选课组合不符合六科结构')
    hl=sum(c['level']=='HL' for c in selected)
    if grade!=10 and hl not in (3,4): raise ValueError('需要3 HL或获批的4 HL；4 HL导入后须登记批准记录')
    if grade==10 and any(c['level']!='NONE' for c in selected): raise ValueError('10年级不分HL/SL')
    d=deepcopy(data); ids={s['id'] for s in d['students']}
    if roster is None:
        roster=[]; n=1
        while len(roster)<count:
            sid=f'ANON-{grade}-{n:04}';n+=1
            if sid not in ids: roster.append((sid,f'匿名学生 {sid}'));ids.add(sid)
    if len(roster)!=count: raise ValueError('人数与名单行数不一致')
    existing={s['id'] for s in d['students']}
    for sid,name in roster:
        if not sid or sid in existing: raise ValueError('学号不能为空或重复：'+sid)
        existing.add(sid)
        d['students'].append(dict(id=sid,name=name or sid,grade=grade,admin_class=admin_class))
        d['choices'].extend(dict(student=sid,course=cid) for cid in choices)
    return normalize(d)


def add_staff_pool(data, code, family, levels, grades, count, max_periods, names=None):
    if not code.strip() or not 1<=count<=30 or not 1<=max_periods<=60: raise ValueError('请填写教师组名称、人数和周课时上限')
    qualified=[c['id'] for c in data['courses'] if c['family']==family and c['level'] in levels and c['grade'] in grades]
    if not qualified and family != 'COMMON': raise ValueError('没有匹配的开课项目')
    d=deepcopy(data)
    names=names or [f'{code} · 教师名额{i+1}' for i in range(count)]
    if len(names)!=count or len(set(names))!=count: raise ValueError('教师名单行数应等于人数且不能重复')
    if any(t['name'] in names for t in d['teachers']): raise ValueError('已有同名教师，请在已有教师的资格表增加课程，避免重复计算师资')
    existing={t['id'] for t in d['teachers']}
    for i,name in enumerate(names):
        tid=f'{code}-{i+1}'
        if tid in existing: raise ValueError('教师组编码已存在，请更换组名或编辑原教师')
        d['teachers'].append(dict(id=tid,name=name,max_periods=max_periods))
        d['qualifications'].extend(dict(teacher=tid,course=c) for c in qualified)
    return normalize(d)


def add_common_courses(data, grade, teacher_map):
    """Explicit teacher IDs, paired class ownership, separate teaching sessions."""
    d=deepcopy(data)
    classes=sorted({s['admin_class'] for s in d['students'] if s['grade']==grade})
    if not classes or '' in classes: raise ValueError('请先录入该年级学生和行政班')
    subjects={'TOK':4 if grade==11 else 3}
    if grade==11: subjects.update({'体育':1,'会考地理':1,'会考历史':1,'会考政治':1,'Career Navigation':1})
    known={t['id'] for t in d['teachers']}
    for name,periods in subjects.items():
        teachers=teacher_map.get(name,[])
        required=(len(classes)+1)//2 if name in ('TOK','Career Navigation') else 1
        if len(teachers)!=required or len(set(teachers))!=required or any(t not in known for t in teachers):
            raise ValueError(f'{name}需指定{required}位已登记教师（TOK/Career每位负责两个班）')
        for i,cl in enumerate(classes):
            aid=f'G{grade}:{name}:{cl}'
            if any(a['id']==aid for a in d['activities']): raise ValueError('公共课已生成，请在公共课程中编辑，勿重复添加')
            d['activities'].append(dict(id=aid,name=name,grade=grade,admin_classes=cl,periods=periods,
                teacher=teachers[i//2] if name in ('TOK','Career Navigation') else teachers[0],
                room_type='操场' if name=='体育' else '普通',room=''))
    return normalize(d)


def choice_template(data, kind='students'):
    wb=Workbook(); ws=wb.active;ws.title='学生选课' if kind=='students' else '选课组合'
    ws.append((['学号','姓名','年级','行政班'] if kind=='students' else ['年级','行政班','人数'])+[f'选课{i}' for i in range(1,7)])
    guide=wb.create_sheet('课程编码');guide.append(['编码','年级','学科','层级','周课时'])
    for c in data['courses']:guide.append([c['id'],c['grade'],c['name'],c['level'],c['periods']])
    note=wb.create_sheet('填写说明');note.append(['每行填写同一学生或同一行政班内的一个六科组合；选课列填写课程编码。'])
    note.append(['只有各学科总人数不足以恢复选课组合；没有名单时按组合人数生成匿名学生。'])
    note.append(['导入追加至当前学期；学号重复会拒绝整次导入。四门HL须另填本学期批准记录。'])
    out=io.BytesIO();wb.save(out);return out.getvalue()


def import_choices(data, raw, filename):
    if not filename.lower().endswith('.xlsx') or len(raw)>10*1024*1024: raise ValueError('请上传不超过10 MB的.xlsx文件')
    from zipfile import ZipFile
    with ZipFile(io.BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist())>60*1024*1024: raise ValueError('文件解压体积过大')
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
    d=deepcopy(data); count=0
    try:
        names=[n for n in ('学生选课','选课组合') if n in wb.sheetnames]
        if len(names)!=1: raise ValueError('请使用单一的“学生选课”或“选课组合”模板')
        ws=wb[names[0]]
        if ws.max_row>20001 or ws.max_column>20: raise ValueError('选课表超出行列限制')
        personal=names[0]=='学生选课'
        expected=(['学号','姓名','年级','行政班'] if personal else ['年级','行政班','人数'])+[f'选课{i}' for i in range(1,7)]
        iterator=ws.iter_rows();header=[c.value for c in next(iterator)]
        if header!=expected: raise ValueError('表头不匹配，请下载当前选课模板')
        for line,cells in enumerate(iterator,2):
            if all(c.value is None for c in cells):continue
            if any(c.data_type=='f' for c in cells):raise ValueError(f'{names[0]} 第{line}行含公式，请粘贴为值')
            row={k:('' if c.value is None else str(c.value).strip()) for k,c in zip(header,cells)}
            try:
                grade=int(row['年级']);n=1 if personal else int(row['人数'])
                if count+n>3000: raise ValueError('本次最多导入3000名学生')
                roster=[(row['学号'],row['姓名'])] if personal else None
                d=add_combination(d,grade,row['行政班'],n,[row[f'选课{i}'] for i in range(1,7)],roster)
                for s in d['students'][-n:]:s['_row']=line
                for ch in d['choices'][-n*6:]:ch['_row']=line
                count+=n
            except (ValueError,KeyError) as exc:raise ValueError(f'{names[0]} 第{line}行：{exc}') from exc
    finally: wb.close()
    if not count:raise ValueError('文件没有学生或组合记录')
    return d,count


def school_demo(per_class=2):
    """Fictitious example with the supplied calendars; labs are explicitly fictional."""
    d=school_dataset();d['config'].update(title='新规则虚构演示',term='DEMO-FLEX',quick_feasible=True)
    for grade in (11,12):
        for cl in range(1,7):
            selections=[f'{f}-{grade}-{l}' for f,l in [('MAT','HL'),('PHY','HL'),('CHE','HL'),('CHA','SL'),('ECO','SL'),('ENB','SL')]]
            if cl%2==0:selections[2]=f'BIO-{grade}-HL'
            if cl%3==0:selections[2]=f'CHE-{grade}-SL';selections[4]=f'ECO-{grade}-HL'
            d=add_combination(d,grade,f'{grade}-{cl}',per_class,selections)
    # Enough explicitly fictitious staff to test joint cross-grade resources.
    used={c['course'] for c in d['choices']}
    for family in CATALOG:
        relevant=[c for c in d['courses'] if c['family']==family and c['id'] in used]
        if relevant:d=add_staff_pool(d,'演示'+family,family,['HL','SL'],[11,12],3,30)
    for grade in (11,12):
        mapping={}
        for name,count in [('TOK',3)]+([('体育',1),('会考地理',1),('会考历史',1),('会考政治',1),('Career Navigation',3)] if grade==11 else []):
            mapping[name]=[]
            for i in range(count):
                tid=f'DEMO-{grade}-{name}-{i+1}';d['teachers'].append(dict(id=tid,name=f'示例{grade}年级{name}教师{i+1}',max_periods=30));mapping[name].append(tid)
        d=add_common_courses(d,grade,mapping)
    for i in range(4):d['rooms'].append(dict(id=f'SIM-LAB{i+1}',name=f'虚构实验室{i+1}',capacity=15,types='普通;物理实验;化学实验;生物实验',preference=2))
    # Freeze section membership for a quick demonstration, never freeze lesson times.
    from collections import defaultdict
    groups=defaultdict(list)
    for row in d['choices']:groups[row['course']].append(row['student'])
    for cid,members in groups.items():
        c=next(c for c in d['courses'] if c['id']==cid);cap=c['capacity']
        for i in range(0,len(members),cap):
            sid=f'{cid}__{i//cap+1}';d['sections'].append(dict(id=sid,course=cid,teacher='',room='',block=0))
            d['enrollments'].extend(dict(student=s,section=sid) for s in members[i:i+cap])
        c['max_sections']=(len(members)+cap-1)//cap
    return normalize(d)
