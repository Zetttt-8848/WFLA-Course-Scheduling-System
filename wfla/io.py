"""Excel/JSON imports and consistent timetable projections/exports."""
import csv
import io
import json
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from .schema import TABLES, empty_dataset, normalize, day_period, active_slots, activity_students


def safe_cell(value):
    if isinstance(value,str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'"+value
    return value


def workbook_bytes(data=None):
    data = data or empty_dataset()
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    for row in [
        ["WFLA 排课模板 v1", "全部示例均为虚构。业务编码按文本输入。"],
        ["格式", "首行是固定字段编码，勿改；鼠标悬停表头看中文说明。每张表支持网页逐行编辑。"],
        ["选课", "每生6行。10年级NONE且无English A；11/12默认3HL，4HL须本学期approved特批。"],
        ["课程族", "English A使用ENGLISH_A；同科HL/SL使用同一family。语言填写zh/en/es等。"],
        ["模式", "灵活并行按学生避免冲突；旧严格区块模式另保留。年级节次、自修与公共课请在对应表配置。"],
        ["班数", "max_sections=0采用容量下限+1的有限候选范围；无解时可显式增加，不代表所有班数都无解。"],
        ["课型", "lab_periods计入periods；双节次数指两节课的次数。实验室多能力用英文分号分隔。"],
        ["导入", "完整导入新版本，不合并、不覆盖旧版本。拒绝公式，不执行宏。"],
    ]:
        ws.append(row)
    config = wb.create_sheet("课表设置")
    config.append(["key","value"])
    for key,value in data["config"].items():
        config.append([key,json.dumps(value,ensure_ascii=False)])
    for key,(label,fields) in TABLES.items():
        ws = wb.create_sheet(label)
        ws.append(list(fields))
        for i,(field,(description,kind,default)) in enumerate(fields.items(),1):
            ws.cell(1,i).comment = Comment(description+f"；默认：{default}","WFLA")
        for row in data.get(key,[]):
            ws.append([safe_cell(row.get(field,"")) for field in fields])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for ws in wb:
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(45,max(16,len(str(col[0].value))+3))
    output=io.BytesIO(); wb.save(output)
    return output.getvalue()


def import_bytes(raw, filename):
    if len(raw)>10*1024*1024:
        raise ValueError("文件不能超过10 MB")
    if filename.lower().endswith(".json"):
        return normalize(json.loads(raw.decode("utf-8-sig")))
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("当前支持.xlsx和.json；CSV用于课表导出")
    # Reject excessive ZIP expansion before openpyxl parses the XML.
    from zipfile import ZipFile
    with ZipFile(io.BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist()) > 60*1024*1024:
            raise ValueError("Excel解压体积过大")
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
    data=empty_dataset()
    if "课表设置" not in wb.sheetnames:
        raise ValueError("缺少课表设置，请下载本系统模板")
    for ws in wb:
        if ws.max_row>20001 or ws.max_column>50:
            raise ValueError(f"{ws.title}超出行列限制")
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError(f"{ws.title} 第{cell.row}行 {cell.column_letter}列含公式，请转换成值")
    for row in list(wb["课表设置"].values)[1:]:
        if row[0]:
            try:
                data["config"][str(row[0])]=json.loads(str(row[1]))
            except (ValueError,TypeError):
                raise ValueError("课表设置格式错误，请保留模板格式或用网页设置")
    for key,(label,fields) in TABLES.items():
        if label not in wb.sheetnames:
            if key in ('calendar','activities','activity_members','study_rules'): continue
            raise ValueError("模板缺少工作表："+label)
        rows=iter(wb[label].values)
        headers=list(next(rows,()))
        legacy_rooms = key == 'rooms' and headers == ['id','name','capacity','types']
        if headers != list(fields) and not legacy_rooms:
            raise ValueError(label+"表头不匹配，请下载当前模板")
        for line,values in enumerate(rows,2):
            if any(v is not None for v in values):
                data[key].append({**dict(zip(headers,values)),"_row":line})
    wb.close()
    return normalize(data)


def timetable_rows(data,result):
    sections={s["id"]:s for s in result.get("sections",[])}
    courses={c["id"]:c for c in data["courses"]}
    rows=[]
    for lesson in result.get("lessons",[]):
        s=sections[lesson["section"]]; c=courses[s["course"]]
        day,period=day_period(data,lesson["start"])
        rows.append({"课次":lesson["id"],"星期":day,"起始节":period,"节数":lesson["duration"],
            "年级":c["grade"],"区块":s["block"],"教学班":s["id"],"课程":c["name"],
            "层级":"不分层" if c["level"]=="NONE" else c["level"],"课型":"实验" if lesson["kind"]=="lab" else "理论/艺术",
            "教师":s["teacher"],"教室":lesson["room"],"人数":len(s["students"]),"学生":";".join(s["students"])})
    for e in data["events"]:
        members=[r["student"] for r in data["event_students"] if r["event"]==e["id"]]
        rows.append({"课次":e["id"],"星期":e["day"],"起始节":e["period"],"节数":e["duration"],"年级":"","区块":"",
            "教学班":"","课程":e["name"],"层级":"","课型":"固定活动","教师":e["teacher"],"教室":e["room"],
            "人数":len(members),"学生":";".join(members)})
    activities = {a['id']:a for a in data.get('activities',[])}
    for lesson in result.get('activities',[]):
        a = activities[lesson['activity']]; members = activity_students(data,a)
        day,period = day_period(data,lesson['start'])
        rows.append({'课次':lesson['id'],'星期':day,'起始节':period,'节数':1,'课程':a['name'],'层级':'',
            '年级':a['grade'],'区块':0,'课型':'公共课','教学班':'COMMON:'+a['id'],'教师':a['teacher'],
            '教室':lesson['room'],'人数':len(members),'学生':';'.join(members)})
    for lesson in result.get('studies',[]):
        day,period = day_period(data,lesson['start']); members = lesson['students']
        rows.append({'课次':lesson['id'],'星期':day,'起始节':period,'节数':1,'课程':'自修','层级':'',
            '年级':'','区块':0,'课型':'自修','教学班':'','教师':'','教室':lesson['room'],'人数':len(members),'学生':';'.join(members)})
    if data['config'].get('scheduling_mode') == 'flexible':
        occupied = set()
        for row in rows:
            for sid in row['学生'].split(';'):
                for offset in range(row['节数']): occupied.add((sid,(row['星期']-1)*data['config']['periods_per_day']+row['起始节']-1+offset))
        for s in data['students']:
            for p in active_slots(data,s['grade']):
                if (s['id'],p) in occupied: continue
                day,period = day_period(data,p)
                rows.append({'课次':f"UNASSIGNED:{s['id']}:{p}",'星期':day,'起始节':period,'节数':1,'课程':'待安排','层级':'',
                    '年级':s['grade'],'区块':0,'课型':'待安排','教学班':'','教师':'','教室':'','人数':1,'学生':s['id']})
    columns = ['课次','星期','起始节','节数','年级','区块','教学班','课程','层级','课型','教师','教室','人数','学生']
    return [{k:r.get(k,'') for k in columns} for r in sorted(rows,key=lambda r:(r["星期"],r["起始节"],r["课次"]))]


def csv_bytes(rows):
    output=io.StringIO(newline="")
    if rows:
        writer=csv.DictWriter(output,fieldnames=list(rows[0])); writer.writeheader()
        writer.writerows({k:safe_cell(v) for k,v in row.items()} for row in rows)
    return output.getvalue().encode("utf-8-sig")


def schedule_xlsx(data,result):
    wb=Workbook(); wb.remove(wb.active)
    rows=timetable_rows(data,result)
    groups={"课次明细":rows}
    for name,field in (("学生视图","学生"),("教师视图","教师"),("教室视图","教室"),("教学班视图","教学班")):
        groups[name]=[{"查看对象":identity,**row} for row in rows for identity in (row[field].split(";") if field=="学生" else [row[field]]) if identity]
    groups["独立审计"]=[{"项目":"硬冲突数量","结果":len(result.get("audit",[]))},{"项目":"求解状态","结果":result["status"]},
                      {"项目":"输入学期","结果":data["config"]["term"]}]
    for name,records in groups.items():
        ws=wb.create_sheet(name)
        if records:
            ws.append(list(records[0]))
            for row in records: ws.append([safe_cell(v) for v in row.values()])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        ws.print_title_rows="1:1"; ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A4; ws.page_setup.fitToWidth=1
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=20
    out=io.BytesIO();wb.save(out);return out.getvalue()
